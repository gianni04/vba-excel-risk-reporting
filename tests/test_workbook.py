"""
test_workbook.py
=================

Verifie que le classeur produit par ``workbook.construire_classeur`` s'ouvre
correctement avec openpyxl et contient les feuilles, les noms definis et les
formules attendus (via ``riskreporting.validate``).
"""

from __future__ import annotations

import openpyxl
import pytest

from riskreporting import validate


def test_classeur_est_ouvrable_avec_openpyxl(classeur_genere):
    classeur = openpyxl.load_workbook(str(classeur_genere), data_only=False)
    assert classeur.sheetnames  # au moins une feuille


def test_toutes_les_feuilles_attendues_sont_presentes(classeur_genere):
    classeur = openpyxl.load_workbook(str(classeur_genere), data_only=False)
    for nom in validate.FEUILLES_ATTENDUES:
        assert nom in classeur.sheetnames


def test_tous_les_noms_definis_attendus_sont_presents(classeur_genere):
    classeur = openpyxl.load_workbook(str(classeur_genere), data_only=False)
    noms_presents = set(classeur.defined_names.keys())
    for nom in validate.NOMS_DEFINIS_ATTENDUS:
        assert nom in noms_presents


def test_table_positions_existe(classeur_genere):
    classeur = openpyxl.load_workbook(str(classeur_genere), data_only=False)
    ws = classeur["Positions"]
    assert "TablePositions" in ws.tables.keys()


def test_validation_complete_est_ok(classeur_genere):
    resultat = validate.valider_classeur(str(classeur_genere))
    assert resultat.ok, resultat.erreurs
    resultat.lever_si_invalide()  # ne doit pas lever


def test_nombre_de_formules_est_substantiel(classeur_genere):
    n = validate.compter_formules(str(classeur_genere))
    assert n > 50


def test_feuille_limites_est_protegee(classeur_genere):
    classeur = openpyxl.load_workbook(str(classeur_genere), data_only=False)
    ws = classeur["Limites"]
    assert ws.protection.sheet


def test_classeur_invalide_leve_une_assertion(tmp_path, jeu_donnees):
    from riskreporting import workbook

    # Classeur volontairement incomplet : feuilles Limites/Graphiques/Documentation
    # absentes (construction manuelle minimale pour verifier la detection d'erreur).
    import xlsxwriter

    chemin = tmp_path / "incomplet.xlsx"
    wb = xlsxwriter.Workbook(str(chemin))
    wb.add_worksheet("Synthese")
    wb.close()

    resultat = validate.valider_classeur(str(chemin))
    assert not resultat.ok
    assert any("Feuille manquante" in e for e in resultat.erreurs)
    with pytest.raises(AssertionError):
        resultat.lever_si_invalide()
